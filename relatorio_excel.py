from io import BytesIO

from openpyxl import Workbook

from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
    Border,
    Side
)

from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo
)

from openpyxl.worksheet.page import PageMargins


MESES = {
    1: "Janeiro",
    2: "Fevereiro",
    3: "Março",
    4: "Abril",
    5: "Maio",
    6: "Junho",
    7: "Julho",
    8: "Agosto",
    9: "Setembro",
    10: "Outubro",
    11: "Novembro",
    12: "Dezembro"
}


# ============================================================
# RELATÓRIO DE EMPRESAS
# ============================================================

def gerar_relatorio_excel(
    empresa_nome,
    mes,
    ano,
    resumo,
    servicos
):

    workbook = Workbook()

    planilha = workbook.active

    planilha.title = "Serviços"

    # ========================================================
    # CORES
    # ========================================================

    azul_escuro = "17365D"
    azul = "2F75B5"
    azul_claro = "EAF2F8"
    cinza_borda = "D8E0EA"
    cinza_texto = "5B6573"
    branco = "FFFFFF"

    preenchimento_titulo = PatternFill(
        "solid",
        fgColor=azul_escuro
    )

    preenchimento_secao = PatternFill(
        "solid",
        fgColor=azul
    )

    preenchimento_cabecalho = PatternFill(
        "solid",
        fgColor=azul_escuro
    )

    preenchimento_card = PatternFill(
        "solid",
        fgColor=azul_claro
    )

    borda = Border(
        left=Side(
            style="thin",
            color=cinza_borda
        ),
        right=Side(
            style="thin",
            color=cinza_borda
        ),
        top=Side(
            style="thin",
            color=cinza_borda
        ),
        bottom=Side(
            style="thin",
            color=cinza_borda
        )
    )

    # ========================================================
    # CONFIGURAÇÃO GERAL
    # ========================================================

    planilha.sheet_view.showGridLines = False
    planilha.sheet_view.zoomScale = 110
    planilha.freeze_panes = "A12"

    # ========================================================
    # LARGURA DAS COLUNAS
    # ========================================================

    planilha.column_dimensions["A"].width = 9
    planilha.column_dimensions["B"].width = 16
    planilha.column_dimensions["C"].width = 68
    planilha.column_dimensions["D"].width = 18

    # ========================================================
    # CABEÇALHO PRINCIPAL
    # ========================================================

    planilha.merge_cells("A1:D1")

    planilha["A1"] = "PL MIL TRANSPORTES"

    planilha["A1"].font = Font(
        bold=True,
        size=20,
        color=branco
    )

    planilha["A1"].fill = preenchimento_titulo

    planilha["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    planilha.row_dimensions[1].height = 38

    # ========================================================
    # SUBTÍTULO
    # ========================================================

    planilha.merge_cells("A2:D2")

    planilha["A2"] = (
        "RELATÓRIO OPERACIONAL DE SERVIÇOS"
    )

    planilha["A2"].font = Font(
        bold=True,
        size=12,
        color=azul_escuro
    )

    planilha["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    planilha.row_dimensions[2].height = 25

    # ========================================================
    # EMPRESA E PERÍODO
    # ========================================================

    planilha.merge_cells("A4:B4")
    planilha.merge_cells("C4:D4")

    planilha["A4"] = (
        f"Empresa: {empresa_nome}"
    )

    planilha["C4"] = (
        f"Período: {MESES[mes]} de {ano}"
    )

    for referencia in ["A4", "C4"]:

        planilha[referencia].font = Font(
            bold=True,
            size=11,
            color=azul_escuro
        )

        planilha[referencia].fill = (
            preenchimento_card
        )

        planilha[referencia].border = borda

        planilha[referencia].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    planilha.row_dimensions[4].height = 28

    # ========================================================
    # RESUMO OPERACIONAL
    # ========================================================

    planilha.merge_cells("A6:D6")

    planilha["A6"] = "RESUMO OPERACIONAL"

    planilha["A6"].font = Font(
        bold=True,
        size=11,
        color=branco
    )

    planilha["A6"].fill = preenchimento_secao

    planilha["A6"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    planilha.row_dimensions[6].height = 24

    indicadores = [
        (
            "A7",
            "SERVIÇOS",
            resumo[0]
        ),
        (
            "B7",
            "KM PERCORRIDOS",
            float(resumo[1])
        ),
        (
            "C7",
            "FUNCIONÁRIOS TRANSPORTADOS",
            resumo[2]
        ),
        (
            "D7",
            "MÊS",
            MESES[mes]
        )
    ]

    for referencia, titulo, valor in indicadores:

        coluna = planilha[referencia].column
        linha = planilha[referencia].row

        descricao = planilha.cell(
            linha,
            coluna,
            titulo
        )

        valor_celula = planilha.cell(
            linha + 1,
            coluna,
            valor
        )

        descricao.font = Font(
            bold=True,
            size=9,
            color=cinza_texto
        )

        descricao.border = borda

        descricao.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        valor_celula.font = Font(
            bold=True,
            size=14,
            color=azul_escuro
        )

        valor_celula.fill = (
            preenchimento_card
        )

        valor_celula.border = borda

        valor_celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        if titulo == "KM PERCORRIDOS":

            valor_celula.number_format = (
                '#,##0.00 "KM"'
            )

    planilha.row_dimensions[7].height = 22
    planilha.row_dimensions[8].height = 30

    # ========================================================
    # SEÇÃO DA TABELA
    # ========================================================

    planilha.merge_cells("A10:D10")

    planilha["A10"] = (
        "REGISTRO DOS SERVIÇOS"
    )

    planilha["A10"].font = Font(
        bold=True,
        size=13,
        color=branco
    )

    planilha["A10"].fill = (
        preenchimento_secao
    )

    planilha["A10"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    planilha.row_dimensions[10].height = 28

    # ========================================================
    # CABEÇALHO DA TABELA
    # ========================================================

    linha_cabecalho = 11

    cabecalhos = [
        "ID",
        "DATA",
        "FUNCIONÁRIOS TRANSPORTADOS",
        "KM REALIZADO"
    ]

    for coluna, cabecalho in enumerate(
        cabecalhos,
        1
    ):

        celula = planilha.cell(
            linha_cabecalho,
            coluna,
            cabecalho
        )

        celula.font = Font(
            bold=True,
            size=11,
            color=branco
        )

        celula.fill = (
            preenchimento_cabecalho
        )

        celula.border = borda

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    planilha.row_dimensions[
        linha_cabecalho
    ].height = 38

    # ========================================================
    # REGISTROS DOS SERVIÇOS
    # ========================================================

    primeira_linha = 12

    linha = primeira_linha

    for servico in servicos:

        planilha.cell(
            linha,
            1,
            servico[0]
        )

        planilha.cell(
            linha,
            2,
            servico[3]
        )

        planilha.cell(
            linha,
            3,
            servico[7]
        )

        planilha.cell(
            linha,
            4,
            float(servico[4])
        )

        planilha.cell(
            linha,
            2
        ).number_format = "DD/MM/YYYY"

        planilha.cell(
            linha,
            4
        ).number_format = (
            '#,##0.00 "KM"'
        )

        for coluna in range(1, 5):

            celula = planilha.cell(
                linha,
                coluna
            )

            celula.font = Font(
                size=11,
                color="202124"
            )

            celula.border = borda

            celula.alignment = Alignment(
                horizontal=(
                    "left"
                    if coluna == 3
                    else "center"
                ),
                vertical="center",
                wrap_text=True
            )

        caracteres = len(
            str(servico[7] or "")
        )

        linhas_estimadas = max(
            1,
            caracteres // 55 + 1
        )

        planilha.row_dimensions[
            linha
        ].height = min(
            72,
            max(
                36,
                linhas_estimadas * 22
            )
        )

        linha += 1

    ultima_linha = linha - 1

    # ========================================================
    # TABELA NATIVA DO EXCEL
    # ========================================================

    if servicos:

        tabela = Table(
            displayName="TabelaServicos",
            ref=(
                f"A{linha_cabecalho}:"
                f"D{ultima_linha}"
            )
        )

        estilo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        tabela.tableStyleInfo = estilo

        planilha.add_table(
            tabela
        )

    else:

        planilha.merge_cells(
            start_row=primeira_linha,
            start_column=1,
            end_row=primeira_linha,
            end_column=4
        )

        celula = planilha.cell(
            primeira_linha,
            1,
            (
                "Nenhum serviço foi "
                "registrado neste período."
            )
        )

        celula.font = Font(
            italic=True,
            size=11,
            color=cinza_texto
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        celula.border = borda

        planilha.row_dimensions[
            primeira_linha
        ].height = 40

    # ========================================================
    # CONFIGURAÇÃO PARA IMPRESSÃO
    # ========================================================

    ultima_linha_impressao = max(
        ultima_linha,
        primeira_linha
    )

    planilha.print_area = (
        f"A1:D{ultima_linha_impressao}"
    )

    planilha.print_title_rows = "11:11"

    planilha.page_setup.orientation = (
        "landscape"
    )

    planilha.page_setup.paperSize = (
        planilha.PAPERSIZE_A4
    )

    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = 0

    planilha.sheet_properties.pageSetUpPr.fitToPage = True

    planilha.page_margins = PageMargins(
        left=0.20,
        right=0.20,
        top=0.45,
        bottom=0.45,
        header=0.20,
        footer=0.20
    )

    planilha.print_options.horizontalCentered = True

    planilha.oddFooter.center.text = (
        "PL Mil Transportes • "
        "Relatório Operacional"
    )

    planilha.oddFooter.right.text = (
        "Página &[Page] de &[Pages]"
    )

    # ========================================================
    # GERAR ARQUIVO
    # ========================================================

    arquivo = BytesIO()

    workbook.save(
        arquivo
    )

    arquivo.seek(0)

    return arquivo


# ============================================================
# RELATÓRIO DE MOTORISTAS
# ============================================================

def gerar_relatorio_motoristas_excel(
    empresa_nome,
    mes,
    ano,
    resumo,
    resumo_motoristas
):

    workbook = Workbook()

    planilha = workbook.active

    planilha.title = "Motoristas"

    # ========================================================
    # CORES
    # ========================================================

    azul_escuro = "17365D"
    azul = "2F75B5"
    azul_claro = "EAF2F8"
    cinza_borda = "D8E0EA"
    cinza_texto = "5B6573"
    branco = "FFFFFF"

    preenchimento_titulo = PatternFill(
        "solid",
        fgColor=azul_escuro
    )

    preenchimento_secao = PatternFill(
        "solid",
        fgColor=azul
    )

    preenchimento_cabecalho = PatternFill(
        "solid",
        fgColor=azul_escuro
    )

    preenchimento_card = PatternFill(
        "solid",
        fgColor=azul_claro
    )

    borda = Border(
        left=Side(
            style="thin",
            color=cinza_borda
        ),
        right=Side(
            style="thin",
            color=cinza_borda
        ),
        top=Side(
            style="thin",
            color=cinza_borda
        ),
        bottom=Side(
            style="thin",
            color=cinza_borda
        )
    )

    # ========================================================
    # RESUMO
    # ========================================================

    quantidade_servicos = resumo[0]

    km_total = float(
        resumo[1]
    )

    funcionarios_transportados = resumo[2]

    valor_total = float(
        resumo[3]
    )

    # ========================================================
    # CONFIGURAÇÃO GERAL
    # ========================================================

    planilha.sheet_view.showGridLines = False
    planilha.sheet_view.zoomScale = 110
    planilha.freeze_panes = "A12"

    # ========================================================
    # LARGURA DAS COLUNAS
    # ========================================================

    planilha.column_dimensions["A"].width = 35
    planilha.column_dimensions["B"].width = 16
    planilha.column_dimensions["C"].width = 18
    planilha.column_dimensions["D"].width = 20
    planilha.column_dimensions["E"].width = 24

    # ========================================================
    # CABEÇALHO PRINCIPAL
    # ========================================================

    planilha.merge_cells("A1:E1")

    planilha["A1"] = "PL MIL TRANSPORTES"

    planilha["A1"].font = Font(
        bold=True,
        size=20,
        color=branco
    )

    planilha["A1"].fill = (
        preenchimento_titulo
    )

    planilha["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    planilha.row_dimensions[1].height = 38

    # ========================================================
    # SUBTÍTULO
    # ========================================================

    planilha.merge_cells("A2:E2")

    planilha["A2"] = (
        "RELATÓRIO DE MOTORISTAS"
    )

    planilha["A2"].font = Font(
        bold=True,
        size=12,
        color=azul_escuro
    )

    planilha["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    planilha.row_dimensions[2].height = 25

    # ========================================================
    # EMPRESA E PERÍODO
    # ========================================================

    planilha.merge_cells("A4:C4")
    planilha.merge_cells("D4:E4")

    planilha["A4"] = (
        f"Empresa: {empresa_nome}"
    )

    planilha["D4"] = (
        f"Período: {MESES[mes]} de {ano}"
    )

    for referencia in ["A4", "D4"]:

        planilha[referencia].font = Font(
            bold=True,
            size=11,
            color=azul_escuro
        )

        planilha[referencia].fill = (
            preenchimento_card
        )

        planilha[referencia].border = borda

        planilha[referencia].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    planilha.row_dimensions[4].height = 28

    # ========================================================
    # RESUMO
    # ========================================================

    planilha.merge_cells("A6:E6")

    planilha["A6"] = (
        "RESUMO DOS MOTORISTAS"
    )

    planilha["A6"].font = Font(
        bold=True,
        size=11,
        color=branco
    )

    planilha["A6"].fill = (
        preenchimento_secao
    )

    planilha["A6"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    planilha.row_dimensions[6].height = 24

    indicadores = [
        (
            "A7",
            "SERVIÇOS",
            quantidade_servicos
        ),
        (
            "B7",
            "KM PERCORRIDOS",
            km_total
        ),
        (
            "C7",
            "FUNCIONÁRIOS TRANSPORTADOS",
            funcionarios_transportados
        ),
        (
            "D7",
            "TOTAL A RECEBER",
            valor_total
        ),
        (
            "E7",
            "MÊS",
            MESES[mes]
        )
    ]

    for referencia, titulo, valor in indicadores:

        coluna = planilha[referencia].column
        linha = planilha[referencia].row

        descricao = planilha.cell(
            linha,
            coluna,
            titulo
        )

        valor_celula = planilha.cell(
            linha + 1,
            coluna,
            valor
        )

        descricao.font = Font(
            bold=True,
            size=9,
            color=cinza_texto
        )

        descricao.border = borda

        descricao.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        valor_celula.font = Font(
            bold=True,
            size=14,
            color=azul_escuro
        )

        valor_celula.fill = (
            preenchimento_card
        )

        valor_celula.border = borda

        valor_celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        if titulo == "KM PERCORRIDOS":

            valor_celula.number_format = (
                '#,##0.00 "KM"'
            )

        if titulo == "TOTAL A RECEBER":

            valor_celula.number_format = (
                'R$ #,##0.00'
            )

    planilha.row_dimensions[7].height = 22
    planilha.row_dimensions[8].height = 30

    # ========================================================
    # SEÇÃO DA TABELA
    # ========================================================

    planilha.merge_cells("A10:E10")

    planilha["A10"] = (
        "RESUMO POR MOTORISTA"
    )

    planilha["A10"].font = Font(
        bold=True,
        size=13,
        color=branco
    )

    planilha["A10"].fill = (
        preenchimento_secao
    )

    planilha["A10"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    planilha.row_dimensions[10].height = 28

    # ========================================================
    # CABEÇALHO
    # ========================================================

    linha_cabecalho = 11

    cabecalhos = [
        "MOTORISTA",
        "SERVIÇOS",
        "KM TOTAL",
        "VALOR POR KM",
        "TOTAL A RECEBER"
    ]

    for coluna, cabecalho in enumerate(
        cabecalhos,
        1
    ):

        celula = planilha.cell(
            linha_cabecalho,
            coluna,
            cabecalho
        )

        celula.font = Font(
            bold=True,
            size=11,
            color=branco
        )

        celula.fill = (
            preenchimento_cabecalho
        )

        celula.border = borda

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    planilha.row_dimensions[
        linha_cabecalho
    ].height = 38

    # ========================================================
    # MOTORISTAS
    # ========================================================

    primeira_linha = 12

    linha = primeira_linha

    for motorista in resumo_motoristas:

        nome = motorista[0]

        quantidade_servicos = motorista[1]

        km = float(
            motorista[2]
        )

        valor_total_motorista = float(
            motorista[3]
        )

        valor_por_km = (
            valor_total_motorista / km
            if km
            else 0
        )

        planilha.cell(
            linha,
            1,
            nome
        )

        planilha.cell(
            linha,
            2,
            quantidade_servicos
        )

        planilha.cell(
            linha,
            3,
            km
        )

        planilha.cell(
            linha,
            4,
            valor_por_km
        )

        planilha.cell(
            linha,
            5,
            valor_total_motorista
        )

        planilha.cell(
            linha,
            3
        ).number_format = (
            '#,##0.00 "KM"'
        )

        planilha.cell(
            linha,
            4
        ).number_format = (
            'R$ #,##0.00'
        )

        planilha.cell(
            linha,
            5
        ).number_format = (
            'R$ #,##0.00'
        )

        for coluna in range(1, 6):

            celula = planilha.cell(
                linha,
                coluna
            )

            celula.font = Font(
                size=11,
                color="202124"
            )

            celula.border = borda

            celula.alignment = Alignment(
                horizontal=(
                    "left"
                    if coluna == 1
                    else "center"
                ),
                vertical="center"
            )

        planilha.row_dimensions[
            linha
        ].height = 28

        linha += 1

    ultima_linha = linha - 1

    # ========================================================
    # TABELA NATIVA DO EXCEL
    # ========================================================

    if resumo_motoristas:

        tabela = Table(
            displayName="TabelaMotoristas",
            ref=(
                f"A{linha_cabecalho}:"
                f"E{ultima_linha}"
            )
        )

        estilo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        tabela.tableStyleInfo = estilo

        planilha.add_table(
            tabela
        )

    else:

        planilha.merge_cells(
            start_row=primeira_linha,
            start_column=1,
            end_row=primeira_linha,
            end_column=5
        )

        celula = planilha.cell(
            primeira_linha,
            1,
            (
                "Nenhum motorista "
                "possui serviços neste período."
            )
        )

        celula.font = Font(
            italic=True,
            size=11,
            color=cinza_texto
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        celula.border = borda

        planilha.row_dimensions[
            primeira_linha
        ].height = 40

    # ========================================================
    # CONFIGURAÇÃO PARA IMPRESSÃO
    # ========================================================

    ultima_linha_impressao = max(
        ultima_linha,
        primeira_linha
    )

    planilha.print_area = (
        f"A1:E{ultima_linha_impressao}"
    )

    planilha.print_title_rows = "11:11"

    planilha.page_setup.orientation = (
        "landscape"
    )

    planilha.page_setup.paperSize = (
        planilha.PAPERSIZE_A4
    )

    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = 0

    planilha.sheet_properties.pageSetUpPr.fitToPage = True

    planilha.page_margins = PageMargins(
        left=0.20,
        right=0.20,
        top=0.45,
        bottom=0.45,
        header=0.20,
        footer=0.20
    )

    planilha.print_options.horizontalCentered = True

    planilha.oddFooter.center.text = (
        "PL Mil Transportes • "
        "Relatório de Motoristas"
    )

    planilha.oddFooter.right.text = (
        "Página &[Page] de &[Pages]"
    )

    # ========================================================
    # GERAR ARQUIVO
    # ========================================================

    arquivo = BytesIO()

    workbook.save(
        arquivo
    )

    arquivo.seek(0)

    return arquivo